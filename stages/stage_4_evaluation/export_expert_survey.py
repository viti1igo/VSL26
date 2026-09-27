"""Generate an expert survey question bank from prescription images and glosses."""

import argparse
import csv
import html
import json
import os
import re
import shutil
from dataclasses import asdict
from pathlib import Path
from typing import Any

from compare_extraction_methods import _mapping_readiness
from llm_extractor import (
    extract_prescription_with_vlm,
    load_project_env,
    vlm_extraction_to_mapper_input,
)
from medicine_mapper import MedicineMapper


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT_DIR = REPO_ROOT / "expert_survey"
DEFAULT_MANIFESTS = [
    REPO_ROOT / "survey_input" / "vaipe_10" / "sample_manifest.csv",
    REPO_ROOT / "survey_input" / "archive_mix_10" / "sample_manifest.csv",
    REPO_ROOT / "survey_input" / "real_world_10" / "sample_manifest.csv",
]
SCORE_PROMPT = (
    "Đánh giá chất lượng gloss này theo thang điểm 1-5 "
    "(1 = không đạt / sai nghiêm trọng / không an toàn; "
    "5 = rất tốt / có thể sử dụng)."
)
COMMENT_PROMPT = "Góp ý chỉnh sửa gloss nếu có (không bắt buộc)."
CHAR_SPLIT_RE = re.compile(r"(?:^|\+\s)([A-Za-zÀ-ỹ])\s\+\s([A-Za-zÀ-ỹ])\s\+\s([A-Za-zÀ-ỹ])")


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))


def _selected_rows(manifest_path: Path) -> list[dict[str, str]]:
    rows = _read_csv(manifest_path)
    manifest_name = manifest_path.parent.name
    selected = []
    for row in rows:
        include = row.get("include_for_30_question_plan", "yes").strip().lower()
        if include in {"backup", "no", "false", "0"}:
            continue
        row = dict(row)
        row["source_group"] = manifest_name
        selected.append(row)

    if manifest_name == "real_world_10" and len(selected) > 10:
        selected = selected[:10]
    return selected


def load_selected_samples(manifest_paths: list[Path]) -> list[dict[str, str]]:
    samples = []
    for manifest_path in manifest_paths:
        if not manifest_path.exists():
            raise FileNotFoundError(f"Missing manifest: {manifest_path}")
        samples.extend(_selected_rows(manifest_path))
    return samples


def _question_id(index: int) -> str:
    return f"Q{index:03d}"


def _copy_question_image(question_id: str, image_path: Path, images_dir: Path) -> Path:
    images_dir.mkdir(parents=True, exist_ok=True)
    output_path = images_dir / f"{question_id}{image_path.suffix.lower()}"
    shutil.copy2(image_path, output_path)
    return output_path


def _flatten_universal_closing(universal_closing: dict[str, Any]) -> list[str]:
    tokens = []
    for value in universal_closing.values():
        if isinstance(value, list):
            for item in value:
                if isinstance(item, list):
                    tokens.extend(str(token) for token in item)
                else:
                    tokens.append(str(item))
    return tokens


def _format_drug_gloss(mapping: dict[str, Any]) -> str:
    drugs = mapping.get("drugs", [])
    lines = []

    section_labels = [
        ("dung_cho", "Dùng cho:"),
        ("dung_nhu_nao", "Dùng như nào:"),
        ("luu_y", "Lưu ý:"),
    ]
    for section_key, label in section_labels:
        section_lines = []
        for drug in drugs:
            tokens = _normalize_display_tokens(
                drug.get("gloss_sections", {}).get(section_key, [])
            )
            if not tokens:
                continue
            if len(drugs) > 1:
                section_lines.append(
                    f"thuốc số {drug.get('drug_number', '')}: " + " + ".join(tokens)
                )
            else:
                section_lines.append(" + ".join(tokens))

        if section_lines:
            lines.append(label)
            lines.extend(section_lines)
            lines.append("")

    return "\n".join(lines).strip()


def _normalize_display_tokens(tokens: Any) -> list[str]:
    if not tokens:
        return []
    if isinstance(tokens, str):
        return [tokens]
    if not isinstance(tokens, list):
        return [str(tokens)]
    return [str(token) for token in tokens if str(token).strip()]


def _format_extracted_fields(parsed_ner: dict[str, Any]) -> str:
    lines = []
    diagnoses = parsed_ner.get("diagnoses", [])
    if diagnoses:
        lines.append("Chẩn đoán: " + " | ".join(str(item) for item in diagnoses))
    for index, drug in enumerate(parsed_ner.get("drugs", []), start=1):
        lines.append(f"Thuốc {index}: {drug.get('name', '')}")
        lines.append(f"  Số lượng: {drug.get('quantity', '')}")
        lines.append(f"  Cách dùng: {drug.get('usage', '')}")
    return "\n".join(lines).strip()


def _gloss_ready(readiness: dict[str, Any]) -> bool:
    drug_count = int(readiness.get("drug_count", 0) or 0)
    return (
        drug_count > 0
        and int(readiness.get("mapped_drug_count", 0) or 0) == drug_count
        and int(readiness.get("usage_complete_count", 0) or 0) == drug_count
    )


def _safe_json_path(question_id: str, results_dir: Path) -> Path:
    results_dir.mkdir(parents=True, exist_ok=True)
    return results_dir / f"{question_id}.json"


def _extract_and_map(
    *,
    question_id: str,
    image_path: Path,
    mapper: MedicineMapper,
    results_dir: Path,
    model: str | None,
    refresh: bool,
) -> dict[str, Any]:
    cache_path = _safe_json_path(question_id, results_dir)
    if cache_path.exists() and not refresh:
        payload = json.loads(cache_path.read_text(encoding="utf-8"))
        if payload.get("raw_vlm_extraction"):
            parsed_ner = vlm_extraction_to_mapper_input(payload["raw_vlm_extraction"])
            mapping = mapper.map_prescription(f"{question_id}_{image_path.stem}", parsed_ner)
            payload["parsed_ner"] = parsed_ner
            payload["mapping"] = asdict(mapping)
            payload["readiness"] = _mapping_readiness(mapping)
            cache_path.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
        return payload

    extraction = extract_prescription_with_vlm(image_path, model=model)
    parsed_ner = vlm_extraction_to_mapper_input(extraction)
    mapping = mapper.map_prescription(f"{question_id}_{image_path.stem}", parsed_ner)
    payload = {
        "question_id": question_id,
        "method": "vlm_direct_extraction",
        "image_path": str(image_path),
        "raw_vlm_extraction": extraction,
        "parsed_ner": parsed_ner,
        "mapping": asdict(mapping),
        "readiness": _mapping_readiness(mapping),
    }
    cache_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def _write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _validate_survey_rows(rows: list[dict[str, Any]]) -> None:
    bad_rows = []
    for row in rows:
        gloss_text = str(row.get("gloss_text", ""))
        match = CHAR_SPLIT_RE.search(gloss_text)
        if match:
            bad_rows.append((row.get("question_id", ""), match.group(0)))

    if bad_rows:
        preview = "; ".join(f"{qid}: {fragment}" for qid, fragment in bad_rows[:5])
        raise ValueError(f"Broken character-split gloss detected: {preview}")


def _write_xlsx(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Survey Questions"
    sheet.append(fieldnames)
    for row in rows:
        sheet.append([row.get(field, "") for field in fieldnames])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    widths = {
        "question_id": 12,
        "image_file": 42,
        "gloss_text": 80,
        "score_prompt": 60,
        "comment_prompt": 45,
        "extracted_fields": 70,
    }
    for index, field in enumerate(fieldnames, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(field, 24)
    for row_index in range(2, len(rows) + 2):
        sheet.row_dimensions[row_index].height = 120
        for cell in sheet[row_index]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def _write_html_preview(path: Path, rows: list[dict[str, Any]], output_dir: Path) -> None:
    cards = []
    for row in rows:
        image_path = Path(row["image_file"])
        try:
            image_src = image_path.relative_to(output_dir)
        except ValueError:
            image_src = image_path
        cards.append(
            f"""
            <section class="question">
              <h2>{html.escape(row["question_id"])}</h2>
              <img src="{html.escape(str(image_src))}" alt="{html.escape(row["question_id"])} prescription image">
              <h3>Generated gloss</h3>
              <pre>{html.escape(row["gloss_text"])}</pre>
              <p><strong>Score:</strong> {html.escape(row["score_prompt"])}</p>
              <p><strong>Optional comment:</strong> {html.escape(row["comment_prompt"])}</p>
            </section>
            """
        )

    path.write_text(
        f"""<!doctype html>
<html lang="vi">
<head>
  <meta charset="utf-8">
  <title>VSL26 Expert Survey Preview</title>
  <style>
    body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; margin: 32px; color: #1f2933; }}
    .question {{ border-top: 1px solid #d0d7de; padding: 24px 0; max-width: 980px; }}
    img {{ max-width: 420px; max-height: 560px; border: 1px solid #d0d7de; }}
    pre {{ white-space: pre-wrap; line-height: 1.5; background: #f6f8fa; padding: 16px; border-radius: 6px; }}
  </style>
</head>
<body>
  <h1>VSL26 Expert Survey Preview</h1>
  <p>Each question should be recreated in the survey form with the prescription image, generated gloss, one required 1-5 score, and one optional comment box.</p>
  {''.join(cards)}
</body>
</html>
""",
        encoding="utf-8",
    )


def build_survey_package(
    *,
    manifest_paths: list[Path],
    output_dir: Path,
    model: str | None,
    auto_enrich: bool,
    refresh: bool,
) -> dict[str, Any]:
    load_project_env()
    if os.getenv("OPENAI_VISION_MODEL") and "OPENAI_MODEL" not in os.environ:
        os.environ["OPENAI_MODEL"] = os.environ["OPENAI_VISION_MODEL"]

    samples = load_selected_samples(manifest_paths)
    if len(samples) != 30:
        print(f"Warning: expected 30 selected images, found {len(samples)}.")

    images_dir = output_dir / "images"
    results_dir = output_dir / "results_json"
    output_dir.mkdir(parents=True, exist_ok=True)

    mapper = MedicineMapper(db_path=str(REPO_ROOT / "vaipe_drugs.db"), auto_enrich=auto_enrich)
    survey_rows = []
    hidden_rows = []

    for index, sample in enumerate(samples, start=1):
        question_id = _question_id(index)
        image_path = (REPO_ROOT / sample["image_file"]).resolve()
        question_image = _copy_question_image(question_id, image_path, images_dir)
        result = _extract_and_map(
            question_id=question_id,
            image_path=image_path,
            mapper=mapper,
            results_dir=results_dir,
            model=model,
            refresh=refresh,
        )

        gloss_text = _format_drug_gloss(result["mapping"])
        extracted_fields = _format_extracted_fields(result["parsed_ner"])
        readiness = result["readiness"]
        gloss_ready = _gloss_ready(readiness)

        survey_rows.append(
            {
                "question_id": question_id,
                "image_file": str(question_image),
                "gloss_text": gloss_text,
                "score_prompt": SCORE_PROMPT,
                "comment_prompt": COMMENT_PROMPT,
                "extracted_fields": extracted_fields,
            }
        )
        hidden_rows.append(
            {
                "question_id": question_id,
                "source_group": sample.get("source_group", ""),
                "sample_id": sample.get("sample_id", ""),
                "prescription_id": sample.get("prescription_id", ""),
                "method": "vlm_direct_extraction",
                "original_image_file": sample.get("image_file", ""),
                "survey_image_file": str(question_image),
                "drug_count": readiness["drug_count"],
                "mapped_drug_count": readiness["mapped_drug_count"],
                "usage_complete_count": readiness["usage_complete_count"],
                "gloss_ready": gloss_ready,
                "missing_vsl_count": readiness["missing_vsl_count"],
                "placeholder_vsl_count": readiness["placeholder_vsl_count"],
                "overall_confidence": result["parsed_ner"].get("vlm_overall_confidence", ""),
                "result_json": str(results_dir / f"{question_id}.json"),
            }
        )
        print(
            f"{question_id}: {sample.get('source_group', '')}/{sample.get('prescription_id', '')} "
            f"drugs={readiness['drug_count']} mapped={readiness['mapped_drug_count']} "
            f"gloss_ready={gloss_ready}"
        )

    survey_fields = [
        "question_id",
        "image_file",
        "gloss_text",
        "score_prompt",
        "comment_prompt",
        "extracted_fields",
    ]
    hidden_fields = [
        "question_id",
        "source_group",
        "sample_id",
        "prescription_id",
        "method",
        "original_image_file",
        "survey_image_file",
        "drug_count",
        "mapped_drug_count",
        "usage_complete_count",
        "gloss_ready",
        "missing_vsl_count",
        "placeholder_vsl_count",
        "overall_confidence",
        "result_json",
    ]

    survey_csv = output_dir / "survey_questions.csv"
    hidden_csv = output_dir / "hidden_mapping.csv"
    survey_xlsx = output_dir / "survey_questions.xlsx"
    response_xlsx = output_dir / "response_analysis_template.xlsx"
    preview_html = output_dir / "survey_preview.html"
    _validate_survey_rows(survey_rows)
    _write_csv(survey_csv, survey_rows, survey_fields)
    _write_csv(hidden_csv, hidden_rows, hidden_fields)
    _write_xlsx(survey_xlsx, survey_rows, survey_fields)
    _write_html_preview(preview_html, survey_rows, output_dir)
    _write_xlsx(
        response_xlsx,
        [
            {
                "question_id": row["question_id"],
                "respondent_code": "",
                "score_1_to_5": "",
                "optional_comment": "",
            }
            for row in survey_rows
        ],
        ["question_id", "respondent_code", "score_1_to_5", "optional_comment"],
    )

    return {
        "survey_csv": str(survey_csv),
        "survey_xlsx": str(survey_xlsx),
        "hidden_csv": str(hidden_csv),
        "response_xlsx": str(response_xlsx),
        "preview_html": str(preview_html),
        "question_count": len(survey_rows),
        "image_count": len(list(images_dir.glob("*"))),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Export VLM-generated glosses for expert survey.")
    parser.add_argument(
        "--manifest",
        action="append",
        default=[],
        help="Input sample manifest CSV. Can be passed multiple times.",
    )
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    parser.add_argument("--model", default=None, help="Override OPENAI_VISION_MODEL.")
    parser.add_argument(
        "--no-auto-enrich",
        action="store_true",
        help="Disable LLM-assisted KB expansion for missing drugs.",
    )
    parser.add_argument("--refresh", action="store_true", help="Ignore cached JSON results.")
    args = parser.parse_args()

    manifest_paths = [Path(item).expanduser() for item in args.manifest] or DEFAULT_MANIFESTS
    summary = build_survey_package(
        manifest_paths=manifest_paths,
        output_dir=Path(args.output_dir).expanduser(),
        model=args.model,
        auto_enrich=not args.no_auto_enrich,
        refresh=args.refresh,
    )
    print("\nSurvey export complete:")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
